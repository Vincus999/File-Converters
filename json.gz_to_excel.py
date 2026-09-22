#!/usr/bin/env python3
"""CNV annotation JSON → XLSX converter.

The input JSON structur:
  { "header": {...}, "positions": [...], "samples": [...] }

Every position becomes a new row in Excel.
The lists (clinvar, decipher, variants) are summarized.
"""
import json
import sys
import glob
from pathlib import Path
import gzip
from typing import cast

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Helper functions
def _join_unique(items, sep=", ", limit=None):
    """Unique value merging, optional limit."""
    seen = []
    for item in items:
        s = str(item)
        if s not in seen:
            seen.append(s)
    if limit:
        seen = seen[:limit]
    return sep.join(seen)


def _safe_get(d, *keys, default=""):
    """Secure nested dict/list access."""
    for key in keys:
        if isinstance(d, dict):
            d = d.get(key, default)
        else:
            return default
    return d if d is not None else default

# GZ uncompress
def _gunzip_json(gz_path):
    """GZ file uncompress and JSON loading."""
    with gzip.open(gz_path, 'rt', encoding='utf-8') as f:
        return json.load(f)

# Position → row conversion
def position_to_row(pos):
    """One position from dict made a row."""
    row = {}

    # Basic data 
    row["chromosome"] = pos.get("chromosome", "")
    row["position"] = pos.get("position", "")
    row["svEnd"] = pos.get("svEnd", "")
    row["svLength"] = pos.get("svLength", "")
    row["id"] = pos.get("id", "")
    row["quality"] = pos.get("quality", "")
    row["cytogeneticBand"] = pos.get("cytogeneticBand", "")
    row["refAllele"] = pos.get("refAllele", "")
    row["altAlleles"] = _join_unique(pos.get("altAlleles", []))
    row["filters"] = _join_unique(pos.get("filters", []))

    ciPos = pos.get("ciPos", [])
    row["ciPos"] = f"{ciPos[0]},{ciPos[1]}" if len(ciPos) >= 2 else ""
    ciEnd = pos.get("ciEnd", [])
    row["ciEnd"] = f"{ciEnd[0]},{ciEnd[1]}" if len(ciEnd) >= 2 else ""

    # First sample data 
    samples = pos.get("samples", [])
    if samples and isinstance(samples[0], dict):
        s = samples[0]
        row["sample_genotype"] = s.get("genotype", "")
        row["sample_copyNumber"] = s.get("copyNumber", "")
        row["sample_simpleNomenclature"] = s.get("simpleNomenclature", "")
        row["sample_binCount"] = s.get("binCount", "")
        row["sample_segmentMean"] = s.get("segmentMean", "")
    else:
        row["sample_genotype"] = ""
        row["sample_copyNumber"] = ""
        row["sample_simpleNomenclature"] = ""
        row["sample_binCount"] = ""
        row["sample_segmentMean"] = ""

    # ClinVar summary 
    clinvar = pos.get("clinvar", [])
    row["clinvar_count"] = len(clinvar)
    all_sig = []
    all_ids = []
    all_review = []
    for cv in clinvar:
        sig = cv.get("significance", [])
        if isinstance(sig, list):
            all_sig.extend(sig)
        else:
            all_sig.append(str(sig))
        all_ids.append(cv.get("id", ""))
        all_review.append(cv.get("reviewStatus", ""))
    row["clinvar_significant"] = _join_unique(all_sig)
    row["clinvar_top_ids"] = _join_unique(all_ids, limit=5)
    row["clinvar_top_review"] = _join_unique(all_review, limit=5)

    # Decipher summary 
    decipher = pos.get("decipher", [])
    row["decipher_count"] = len(decipher)
    max_del = 0.0
    max_dup = 0.0
    for dc in decipher:
        max_del = max(max_del, dc.get("deletionFrequency", 0))
        max_dup = max(max_dup, dc.get("duplicationFrequency", 0))
    row["decipher_maxDelFreq"] = max_del if max_del > 0 else ""
    row["decipher_maxDupFreq"] = max_dup if max_dup > 0 else ""

    # Variant + gene data
    variants = pos.get("variants", [])
    if variants and isinstance(variants[0], dict):
        v = variants[0]
        row["variant_type"] = v.get("variantType", "")
        row["variant_nomenclature"] = v.get("simpleNomenclature", "")

        # Genes, consequence, impact from transcripts
        genes = []
        consequences = []
        impacts = []
        for tr in v.get("transcripts", []):
            g = tr.get("hgnc", "")
            if g:
                genes.append(g)
            cons = tr.get("consequence", [])
            if isinstance(cons, list):
                consequences.extend(cons)
            imp = tr.get("impact", "")
            if imp:
                impacts.append(imp)
        row["genes"] = _join_unique(genes)
        row["consequences"] = _join_unique(consequences)
        row["impacts"] = _join_unique(impacts)
    else:
        row["variant_type"] = ""
        row["variant_nomenclature"] = ""
        row["genes"] = ""
        row["consequences"] = ""
        row["impacts"] = ""

    return row


# Column order 
COLUMN_ORDER = [
    "chromosome", "position", "svEnd", "svLength", "id", "quality",
    "cytogeneticBand", "refAllele", "altAlleles", "filters",
    "ciPos", "ciEnd",
    "sample_genotype", "sample_copyNumber", "sample_simpleNomenclature",
    "sample_binCount", "sample_segmentMean",
    "clinvar_count", "clinvar_significant", "clinvar_top_ids",
    "clinvar_top_review",
    "decipher_count", "decipher_maxDelFreq", "decipher_maxDupFreq",
    "variant_type", "variant_nomenclature", "genes", "consequences",
    "impacts",
]


# Main conversion function
def convert(json_path: str, xlsx_path: str) -> None:
    json_file = Path(json_path)
    if not json_file.exists():
        raise FileNotFoundError(f"Nem található a fájl: {json_path}")

    with json_file.open("r", encoding="utf-8") as f:
        if json_file.suffix == ".gz":
            data = _gunzip_json(json_file)
        else:
            data = json.load(f)

    positions = data.get("positions", []) if isinstance(data, dict) else []

    # Generate rows
    rows = [position_to_row(pos) for pos in positions]

    # Headers
    headers = COLUMN_ORDER[:]

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "CNV Adatok"

    # Header row
    ws.append(headers)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    # Data
    for row_dict in rows:
        ws.append([row_dict.get(h, "") for h in headers])

    # Column width automatic estimation
    for i, header in enumerate(headers, start=1):
        max_len = len(str(header))
        for row_dict in rows:
            val = row_dict.get(header, "")
            max_len = max(max_len, len(str(val)))
        ws.column_dimensions[get_column_letter(i)].width = min(max_len + 2, 50)

    # Pinning headers
    ws.freeze_panes = "A2"

    # filters for headers
    ws.auto_filter.ref = ws.dimensions

    wb.save(xlsx_path)
    if not rows:
        print(f"WARNING: There is no position data in the JSON file. Empty Excel (headers only) is generated: {xlsx_path}")
    else:
        print(f"Complete! {len(rows)} CNV position saved to: {xlsx_path}")


# Batch conversion
def convert_batch(json_files):
    """More JSON file's conversion to XLSX to same folder."""
    success = 0
    empty = 0
    errors = 0

    for json_path in json_files:
        xlsx_path = str(json_path).rsplit(".", 1)[0] + ".xlsx"
        try:
            convert(str(json_path), xlsx_path)
            success += 1
        except Exception as e:
            print(f"HIBA: {json_path.name}: {e}")
            errors += 1

    print(f"\nÖsszesítés: {success} sikeres, {errors} hiba")


# GZ tömörített fájlok kicsomagolása, törlése és konvertálása
def decompress_gz(gz_path, delete_gz=True):
    """GZ fájl kicsomagolása JSON-ba, majd a GZ törlése."""
    json_path = str(gz_path).rsplit(".gz", 1)[0]
    with gzip.open(gz_path, 'rb') as f_in:
        with open(json_path, 'wb') as f_out:
            data = cast(bytes, f_in.read())
            f_out.write(data)
    if delete_gz:
        gz_path.unlink()
    return json_path


def convert_gz_batch(gz_files):
    """Több GZ fájl kicsomagolása + konvertálása XLSX-be."""
    decompressed = 0
    converted = 0
    errors = 0

    for gz_path in gz_files:
        try:
            json_path = decompress_gz(gz_path)
            xlsx_path = str(Path(json_path).with_suffix(".xlsx"))
            convert(json_path, xlsx_path)
            decompressed += 1
            converted += 1
        except Exception as e:
            print(f"HIBA: {gz_path.name}: {e}")
            errors += 1

    print(f"\nÖsszesítés: {decompressed} kicsomagolva, {converted} konvertálva, {errors} hiba")



# Parancssori belépés
def main():
    if len(sys.argv) == 1:
        json_path = input("JSON fájl elérési útja (vagy * glob): ").strip()
        xlsx_path = input("Kimeneti XLSX fájl neve (üreshagyva = automatikus): ").strip()
        if xlsx_path:
            convert(json_path, xlsx_path)
        else:
            files = [Path(f) for f in sorted(glob.glob(str(Path(json_path).expanduser())))]
            if not files:
                print(f"Nem található fájl a mintára: {json_path}")
                sys.exit(1)
            json_files = [f for f in files if not f.suffix == ".gz"]
            gz_files = [f for f in files if f.suffix == ".gz"]
            if json_files:
                convert_batch(json_files)
            if gz_files:
                convert_gz_batch(gz_files)
    elif len(sys.argv) == 2:
        arg = sys.argv[1]
        p = Path(arg)
        if p.is_dir():
            json_files = sorted(p.glob("*.json"))
            gz_files = sorted(p.glob("*.gz"))
            if not json_files and not gz_files:
                print(f"Nincs .json vagy .gz fájl ebben a könyvtárban: {arg}")
                sys.exit(1)
            if json_files:
                print(f"{len(json_files)} JSON fájl találva itt: {arg}")
                convert_batch(json_files)
            if gz_files:
                print(f"{len(gz_files)} GZ fájl találva itt: {arg}")
                convert_gz_batch(gz_files)
        elif "*" in arg:
            files = [Path(f) for f in sorted(glob.glob(str(Path(arg).expanduser())))]
            if not files:
                print(f"Nem található fájl a mintára: {arg}")
                sys.exit(1)
            json_files = [f for f in files if not f.suffix == ".gz"]
            gz_files = [f for f in files if f.suffix == ".gz"]
            if json_files:
                print(f"{len(json_files)} JSON fájl találva a mintára: {arg}")
                convert_batch(json_files)
            if gz_files:
                print(f"{len(gz_files)} GZ fájl találva a mintára: {arg}")
                convert_gz_batch(gz_files)
        else:
            if p.suffix == ".gz":
                gz_files = [p]
                convert_gz_batch(gz_files)
            else:
                xlsx_path = str(p.with_suffix(".xlsx"))
                convert(arg, xlsx_path)
    elif len(sys.argv) == 3:
        convert(sys.argv[1], sys.argv[2])
    else:
        print("Használat:")
        print("  python json_to_excel.py fájl.json                # egy fájl → .xlsx")
        print("  python json_to_excel.py fájl.json.gz             # egy GZ → kicsomagol + .xlsx")
        print("  python json_to_excel.py fájl.json kimenet.xlsx    # egy fájl → megadott név")
        print("  python json_to_excel.py '*.json'                  # mindet konvertálja")
        print("  python json_to_excel.py '*.gz'                    # GZ-ek kicsomagolása + konvertálás")
        print("  python json_to_excel.py /útvonal/                 # könyvtár összes .json/.gz")
        sys.exit(1)


if __name__ == "__main__":
    main()
